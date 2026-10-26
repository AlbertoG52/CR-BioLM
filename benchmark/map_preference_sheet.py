"""Blind map choice: an expert picks the most accurate of four maps for each species.

For each species the expert sees its entry from the Manual and four maps drawn the same
way, labelled only A to D in an order drawn at random per species:

  ours       the Manual's range plus the regions holding at least five records
  manual     the Manual's stated range alone
  mistral    the range Mistral Large 4 describes when asked in plain Spanish
  gemma      the range Gemma 4 31B describes

The maps show the range only. Collection records are deliberately left off: the record
comparison already measures how many records each map captures, and an expert shown the
points would be judging by them rather than by what is known of the plant.

Where two maps are the same (ours and the Manual's, when no record adds a region) both
appear, and the expert may mark more than one.

How the species are chosen, so it can be reported as it was done: from the 300-species
map sample, two sets of 20, drawn at random within each tertile of worldwide record
count (7, 7 and 6), with 5 species placed in both sets so that agreement between the two
experts can be measured. Seed recorded.

Run:  python -m benchmark.map_preference_sheet
Writes map_choice_<expert>.xlsx and map_choice_key.json
"""
from __future__ import annotations

import io
import json
import random
import tempfile
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib.colors import ListedColormap
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from . import mapcompare as M

DATA = Path(__file__).resolve().parent / "data"
SEED = "20261023-map-choice"
EXPERTS = ("Armando", "Nelson")
PER_SET, SHARED = 20, 5
SOURCES = ("ours", "manual", "paid", "open")
FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E46")
FILL = PatternFill("solid", fgColor="FFF2CC")
GREY = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
LAND, FILLC, INK = "#e3e6e1", "#2a7f74", "#1f2a28"


def select() -> dict:
    sample = json.loads((DATA / "maps_sample.json").read_text(encoding="utf-8"))["sample"]
    rng = random.Random(SEED)
    sets = {e: [] for e in EXPERTS}
    shared_quota, own_quota = {"G1": 2, "G2": 2, "G3": 1}, {"G1": 5, "G2": 5, "G3": 5}
    for g in ("G1", "G2", "G3"):
        pool = sorted(r["species"] for r in sample if r["stratum"] == g)
        rng.shuffle(pool)
        shared = pool[:shared_quota[g]]
        rest = pool[shared_quota[g]:]
        n = own_quota[g]
        for i, e in enumerate(EXPERTS):
            sets[e] += shared + rest[i * n:(i + 1) * n]
    for e in EXPERTS:
        rng.shuffle(sets[e])
    return sets


def _ingredients(cur, species: str) -> dict:
    from mpcr_rag.schema import Ficha
    rows = [json.loads(l) for l in io.open(DATA / "mapcompare_scores_maps.jsonl", encoding="utf-8")]
    by = {r["source"]: r["ingredients"] for r in rows if r["species"] == species}
    cur.execute("SELECT ficha FROM mpcr.fichas WHERE species = %s", (species,))
    f = Ficha(**cur.fetchone()[0])
    cur.execute("""SELECT o.region, o.vertiente, o.elev_dem FROM mpcr.occurrences o
                   JOIN mpcr.species_evidence e ON e.taxon_key = o.species_key
                   WHERE e.species = %s""", (species,))
    recs = [(a, None if v is None else ("Caribe" if "arib" in v else "Pacífico"), el)
            for a, v, el in cur.fetchall()]
    manual = M.from_ficha(f)
    return {"ficha": f, "ours": M.with_records(manual, recs), "manual": manual,
            "paid": by.get("paid"), "open": by.get("open")}


def _panel_png(order: list[str], ings: dict, path: Path) -> None:
    M._grid()
    land, b = M._GEO["land"], M._GEO["bounds"]
    extent = (b.left, b.right, b.bottom, b.top)
    fig, axes = plt.subplots(1, 4, figsize=(14, 4.4), dpi=110)
    for ax, letter, src in zip(axes, "ABCD", order):
        ax.imshow(np.where(land, 1.0, np.nan), extent=extent,
                  cmap=ListedColormap([LAND]), interpolation="nearest")
        ing = ings.get(src)
        mk = M.mask(ing) if ing else None
        if mk is not None:
            full = np.full(land.shape, np.nan)
            full[land] = np.where(mk, 1.0, np.nan)
            ax.imshow(full, extent=extent, cmap=ListedColormap([FILLC]), alpha=0.9,
                      interpolation="nearest")
        ax.set_title(f"Mapa {letter}", fontsize=13, color=INK, fontweight="bold")
        ax.set_xticks([])
        ax.set_yticks([])
        for s in ax.spines.values():
            s.set_visible(False)
    fig.subplots_adjust(left=0.01, right=0.99, top=0.9, bottom=0.02, wspace=0.03)
    fig.savefig(path, facecolor="white")
    plt.close(fig)


def _cell(ws, ref, value, *, bold=False, fill=None, color="000000", size=10):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    c.alignment = WRAP
    if fill:
        c.fill = fill
    return c


INSTRUCTIONS = [
    ("Elección de mapas de distribución", True, 14),
    ("", False, 10),
    ("Para cada especie hay cuatro mapas, A, B, C y D, dibujados de la misma forma. El área "
     "verde es donde cada mapa dice que la especie puede encontrarse en Costa Rica. Los mapas "
     "salen de fuentes distintas y el orden cambia de una especie a otra; no se indica cuál "
     "es cuál.", False, 10),
    ("", False, 10),
    ("Qué le pedimos, por especie (solo las celdas amarillas):", True, 11),
    ("1. ¿Cuál mapa representa mejor dónde se encuentra la especie? Elija una letra. Si dos o "
     "más le parecen igual de buenos, elija «Empate» y escriba las letras en el comentario.",
     False, 10),
    ("2. ¿Cuáles mapas considera aceptables? Escriba las letras (por ejemplo «A, C»), o "
     "«Ninguno».", False, 10),
    ("3. Un comentario, si quiere: qué le falta o le sobra al mejor mapa.", False, 10),
    ("", False, 10),
    ("Junto a cada especie está su entrada del Manual de Plantas de Costa Rica, por si quiere "
     "refrescar la especie. A veces dos mapas son idénticos; es normal.", False, 10),
    ("Son 20 especies; calculamos entre 20 y 30 minutos. La primera fila, en gris, es un "
     "ejemplo ya llenado.", False, 10),
]


def build() -> None:
    from mpcr_rag.store import pg_store
    sets = select()
    rng = random.Random(SEED + "-order")
    key: dict = {"seed": SEED, "rule": "two sets of 20 from the 300-species map sample, "
                 "random within worldwide-record tertiles (7, 7, 6), 5 species in both sets; "
                 "map order random per species and expert", "sets": {}}
    conn = pg_store.connect()
    with conn.cursor() as cur, tempfile.TemporaryDirectory() as tmp:
        for expert, species in sets.items():
            wb = Workbook()
            ins = wb.active
            ins.title = "Instrucciones"
            ins.column_dimensions["A"].width = 110
            for i, (t, bold, size) in enumerate(INSTRUCTIONS, 1):
                _cell(ins, f"A{i}", t, bold=bold, size=size)
                if len(t) > 110:
                    ins.row_dimensions[i].height = 15 * (len(t) // 100 + 1)
            ws = wb.create_sheet("Mapas")
            heads = [("N.º", 5), ("Especie y entrada del Manual", 46), ("Mapas", 150),
                     ("¿Cuál representa mejor dónde se encuentra?", 16),
                     ("¿Cuáles son aceptables?", 16), ("Comentario", 34)]
            for col, (t, w) in enumerate(heads, 1):
                c = ws.cell(row=1, column=col, value=t)
                c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
                c.fill = HEAD
                c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                ws.column_dimensions[c.column_letter].width = w
            ws.row_dimensions[1].height = 40
            ws.freeze_panes = "C2"
            # Example row, in grey, filled in.
            for col, v in enumerate(["Ej.", "Especie de ejemplo", "(aquí van los cuatro mapas)",
                                     "C", "B, C", "EJEMPLO. C es el mejor; B también es "
                                     "aceptable aunque le falta la Península de Osa."], 1):
                c = ws.cell(row=2, column=col, value=v)
                c.font = Font(name=FONT, size=10, italic=True)
                c.fill = GREY
                c.alignment = WRAP
                c.border = BOX
            ws.row_dimensions[2].height = 40
            rows_key = []
            for i, sp in enumerate(species, 1):
                ings = _ingredients(cur, sp)
                order = list(SOURCES)
                rng.shuffle(order)
                f = ings["ficha"]
                png = Path(tmp) / f"{expert}_{i}.png"
                _panel_png(order, ings, png)
                r = i + 2
                text = (f"{sp}\n{f.family}. Manual de Plantas de Costa Rica, Tomo {f.volume}, "
                        f"p. {f.pages}.\n\n{f.distribution_paragraph}")
                for col, v in enumerate([i, text, "", "", "", ""], 1):
                    c = ws.cell(row=r, column=col, value=v)
                    c.font = Font(name=FONT, size=10, italic=(col == 2 and False))
                    c.alignment = WRAP
                    c.border = BOX
                    if col >= 4:
                        c.fill = FILL
                ws.cell(row=r, column=2).font = Font(name=FONT, size=10)
                img = XLImage(str(png))
                img.width, img.height = 1050, 330
                ws.add_image(img, f"C{r}")
                ws.row_dimensions[r].height = 255
                rows_key.append({"n": i, "species": sp,
                                 "letters": dict(zip("ABCD", order)),
                                 "identical_ours_manual": ings["ours"] == ings["manual"]})
            dv = DataValidation(type="list", formula1='"A,B,C,D,Empate,No sé"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"D3:D{len(species) + 2}")
            out = DATA / f"map_choice_{expert}.xlsx"
            wb.save(out)
            key["sets"][expert] = rows_key
            print(f"{expert}: {len(species)} species -> {out.name} "
                  f"({out.stat().st_size // 1024} KB); ours identical to the Manual's in "
                  f"{sum(r['identical_ours_manual'] for r in rows_key)}")
    conn.close()
    shared = set(sets[EXPERTS[0]]) & set(sets[EXPERTS[1]])
    key["shared_species"] = sorted(shared)
    (DATA / "map_choice_key.json").write_text(json.dumps(key, ensure_ascii=False, indent=1),
                                              encoding="utf-8")
    print(f"species in both sets: {len(shared)}")


if __name__ == "__main__":
    build()
