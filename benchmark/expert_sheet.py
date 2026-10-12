"""The workbook an expert fills in: thirty questions, two answers each, and the maps.

For each question the expert sees what the record supports and two answers to it, one
written from the evidence pack and one by the same model asked freely. Which is which
is not shown, and their order is randomised per row. The expert says which they would
trust, whether either contains a botanical error, and whether our expected answer is
right. A second sheet carries the Manual entry and the distribution map of each
species, with one question about the map.

How the thirty are chosen, so it can be reported as it was done:

  frame      the questions of tranche 1 (27 species never run before, 3 per cell of the
             sampling grid) that both arms answered
  strata     the twelve question types
  allocation 3 from each of the six types whose key rests on a rule or on silence (the
             three park types, cultivation, co-occurrence, edible or medicinal use), 2
             from each of the other six. That is where an expert's judgement is worth
             most, and the allocation is fixed here, before any verdict is looked at
  draw       simple random sample without replacement within the type, by a recorded
             seed. Verdicts are not consulted

Which answer is which, and what the checker made of each, go to a key file that is not
sent with the workbook.

Run:  python -m benchmark.expert_sheet
Writes expert_review_tramo1.xlsx and expert_review_key.jsonl
"""
from __future__ import annotations

import io
import json
import random
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PILImage

from .validation_sheet import key_text

DATA = Path(__file__).resolve().parent / "data"
SEED = 20261011
GROUNDED, FREE = "D", "A1F"          # same generator, with and without evidence
THREE = ("park_confirmed", "park_possible", "park_unlikely", "cultivation",
         "co_occurrence", "outside_record")
TWO = ("elevation", "vertiente", "flowering", "growth_form", "endemism", "light")
TIPO = {"elevation": "Altitud", "vertiente": "Vertiente", "flowering": "Floración",
        "growth_form": "Forma de crecimiento", "endemism": "Endemismo",
        "park_confirmed": "Área protegida", "park_possible": "Área protegida",
        "park_unlikely": "Área protegida", "cultivation": "Cultivo",
        "co_occurrence": "Especies acompañantes", "outside_record": "Usos",
        "light": "Luz"}
MESES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "set", "oct", "nov", "dic"]

FONT = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E46")
FILL = PatternFill("solid", fgColor="FFF2CC")      # cells the expert fills in
GREY = PatternFill("solid", fgColor="E7E6E6")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def _jsonl(name: str) -> list[dict]:
    return [json.loads(l) for l in io.open(DATA / name, encoding="utf-8") if l.strip()]


def select() -> list[dict]:
    species = set(json.loads((DATA / "species_tranche1.json").read_text(encoding="utf-8")))
    answers = {(a["qid"], a["arm"]): a for a in _jsonl("answers_main.jsonl")
               if not a.get("error") and a.get("answer_text")}
    pool: dict[str, list[dict]] = {}
    for q in _jsonl("questions_main.jsonl"):
        if (q["species"] in species and (q["qid"], GROUNDED) in answers
                and (q["qid"], FREE) in answers):
            pool.setdefault(q["qtype"], []).append(q)
    rng = random.Random(SEED)
    picked = []
    for qtype in THREE + TWO:
        qs = sorted(pool.get(qtype, []), key=lambda q: q["qid"])
        for q in rng.sample(qs, min(3 if qtype in THREE else 2, len(qs))):
            flip = rng.random() < 0.5
            a, b = (FREE, GROUNDED) if flip else (GROUNDED, FREE)
            picked.append({"q": q, "A": a, "B": b,
                           "text_A": answers[(q["qid"], a)]["answer_text"],
                           "text_B": answers[(q["qid"], b)]["answer_text"]})
    rng.shuffle(picked)
    return picked


def _cell(ws, ref, value, *, bold=False, fill=None, color="000000", size=10):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, bold=bold, color=color, size=size)
    c.alignment = WRAP
    if fill:
        c.fill = fill
    return c


def _instructions(ws) -> None:
    ws.title = "Instrucciones"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Revisión experta de respuestas sobre la flora de Costa Rica", True, 14),
        ("", False, 10),
        ("Gracias por su tiempo. Son 30 preguntas y deberían tomar entre 45 minutos y una "
         "hora. No hace falta consultar nada: nos interesa su criterio.", False, 10),
        ("", False, 10),
        ("Qué va a ver", True, 11),
        ("Hoja «Revisión»: una fila por pregunta. Cada fila trae la pregunta, lo que "
         "respalda el registro (el Manual de Plantas de Costa Rica más los registros de "
         "GBIF) y dos respuestas, A y B, escritas por un sistema automático. No se indica "
         "cómo se generó cada una, y el orden cambia de fila en fila.", False, 10),
        ("Hoja «Fichas y mapas»: para cada especie, el texto del Manual tal como lo tenemos "
         "y el mapa de distribución que genera el sistema. El nombre de la especie en la "
         "hoja «Revisión» lleva a su ficha.", False, 10),
        ("", False, 10),
        ("Qué le pedimos", True, 11),
        ("Solo se llenan las celdas amarillas. Las cuatro primeras tienen lista "
         "desplegable.", False, 10),
        ("1. ¿Cuál respuesta le merece más confianza? A, B o Iguales.", False, 10),
        ("2. ¿La respuesta A contiene algún error botánico? Sí, No o No sé. Lo mismo para "
         "la B. Si responde Sí, anote cuál en el comentario.", False, 10),
        ("3. ¿Lo que respalda el registro es correcto? Sí; No; El registro está incompleto "
         "(la especie se conoce de más lugares, altitudes o meses de los que indica); o No "
         "sé.", False, 10),
        ("4. Si la respuesta esperada no es correcta o no existe (por ejemplo, en las "
         "preguntas de luz), ¿cuál sería la correcta?", False, 10),
        ("5. Comentario libre: cualquier observación, por breve que sea, nos sirve.",
         False, 10),
        ("En «Fichas y mapas»: si el mapa es consistente con lo que usted conoce de la "
         "especie (Sí, En parte, No) y un comentario.", False, 10),
        ("", False, 10),
        ("Aclaraciones", True, 11),
        ("«Lo que respalda el registro» no pretende ser la verdad de campo, sino lo que "
         "dicen el Manual y los registros. Si usted sabe que la realidad es otra, eso es "
         "justamente lo que queremos saber.", False, 10),
        ("En las preguntas de área protegida, «posible» significa que no hay registros "
         "dentro del área pero coinciden la región botánica y la altitud; «improbable», que "
         "no hay registros ni coincide ninguna de las dos.", False, 10),
        ("La primera fila de la hoja «Revisión», en gris, es un ejemplo ya llenado. No "
         "hay que evaluarla.", False, 10),
    ]
    for i, (text, bold, size) in enumerate(lines, 1):
        _cell(ws, f"A{i}", text, bold=bold, size=size)
        if len(text) > 110:
            ws.row_dimensions[i].height = 15 * (len(text) // 105 + 1)


HEADERS = [("N.º", 5), ("Especie", 20), ("Tema", 14), ("Pregunta", 34),
           ("Lo que respalda el registro", 40), ("Respuesta A", 60), ("Respuesta B", 60),
           ("¿Cuál le merece más confianza?", 16), ("¿A tiene un error botánico?", 15),
           ("¿B tiene un error botánico?", 15), ("¿Lo que respalda el registro es correcto?", 20),
           ("Si no lo es, ¿cuál sería la respuesta correcta?", 34), ("Comentario", 40)]

EXAMPLE = {
    "species": "Anthurium ravenii", "tema": "Área protegida",
    "question": "Si visito el Parque Nacional Rincón de la Vieja, ¿puedo encontrarme con "
                "Anthurium ravenii?",
    "key": "Clave: posible. Registros dentro del área: 0. El área coincide con las regiones "
           "botánicas y con la altitud de la especie.",
    "A": "Sí, es posible encontrarla, ya que es endémica de Costa Rica y se distribuye "
         "en los bosques nubosos de la Cordillera de Guanacaste, donde se ubica el parque.",
    "B": "Es posible, aunque no existen registros de la especie dentro del parque: el "
         "área coincide con sus regiones botánicas y con su rango altitudinal, y la "
         "ausencia de registros no es evidencia de ausencia.",
    "fill": ["B", "Sí", "No", "Sí", "",
             "EJEMPLO. La A dice que es endémica y el Manual indica Honduras a Colombia."],
}


def _review(ws, items: list[dict], anchors: dict[str, int]) -> None:
    ws.title = "Revisión"
    for col, (title, width) in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BOX
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[1].height = 44
    ws.freeze_panes = "D2"

    def write(row, n, species, tema, question, key, a, b, filled=None, grey=False):
        vals = [n, species, tema, question, key, a, b] + (filled or [""] * 6)
        for col, v in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = Font(name=FONT, size=10, italic=grey)
            c.alignment = WRAP
            c.border = BOX
            c.fill = GREY if grey else (FILL if col >= 8 else PatternFill())
        longest = max(len(str(a)), len(str(b)), len(str(key)) * 1.5)
        ws.row_dimensions[row].height = min(400, max(75, 13.5 * (longest / 62 + 1)))

    write(2, "Ej.", EXAMPLE["species"], EXAMPLE["tema"], EXAMPLE["question"], EXAMPLE["key"],
          EXAMPLE["A"], EXAMPLE["B"], EXAMPLE["fill"], grey=True)
    for i, it in enumerate(items, 1):
        q = it["q"]
        write(i + 2, i, q["species"], TIPO[q["qtype"]], q["question"], key_text(q),
              it["text_A"], it["text_B"])
        link = ws.cell(row=i + 2, column=2)
        link.hyperlink = f"#'Fichas y mapas'!A{anchors[q['species']]}"
        link.font = Font(name=FONT, size=10, color="0563C1", underline="single",
                         italic=True)

    last = len(items) + 2
    for col, options in (("H", "A,B,Iguales"), ("I", "Sí,No,No sé"), ("J", "Sí,No,No sé"),
                         ("K", "Sí,No,El registro está incompleto,No sé")):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        dv.error = "Elija una opción de la lista."
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}{last}")


def _fichas(ws, species: list[str], tmp: Path) -> dict[str, int]:
    from mpcr_rag.mcp import server as S
    from mpcr_rag.schema import Ficha
    from mpcr_rag.store import pg_store
    from .regions import effective

    ws.title = "Fichas y mapas"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    render = getattr(S.render_species_map, "fn", S.render_species_map)
    dv = DataValidation(type="list", formula1='"Sí,En parte,No"', allow_blank=True)
    ws.add_data_validation(dv)
    anchors, row = {}, 1
    conn = pg_store.connect()
    with conn.cursor() as cur:
        for sp in species:
            cur.execute("SELECT ficha FROM mpcr.fichas WHERE species = %s", (sp,))
            f = Ficha(**cur.fetchone()[0])
            cur.execute("SELECT n_records FROM mpcr.species_evidence WHERE species = %s", (sp,))
            got = cur.fetchone()
            n_snapshot = got[0] if got else 0
            reg = effective(sp, f.regions)
            anchors[sp] = row
            _cell(ws, f"A{row}", sp, bold=True, size=12, color="FFFFFF", fill=HEAD)
            _cell(ws, f"B{row}", f"{f.family}. Manual de Plantas de Costa Rica, Tomo "
                                 f"{f.volume}, p. {f.pages}.", color="FFFFFF", fill=HEAD)
            elev = f"{f.elev_min}-{f.elev_max} m"
            if f.elev_min is None:
                # No range was parsed for this entry (34 of 6,946). Printing "None-None"
                # would read as a fact; say what happened, which also explains why the
                # map carries no elevation band.
                elev = ("no extraída de nuestra copia del Manual; por eso el mapa no "
                        "muestra la franja altitudinal")
            elif f.elev_outlier_min is not None or f.elev_outlier_max is not None:
                lo = f.elev_outlier_min if f.elev_outlier_min is not None else f.elev_min
                hi = f.elev_outlier_max if f.elev_outlier_max is not None else f.elev_max
                elev += f" (con registros atípicos {lo}-{hi} m)"
            facts = [
                ("Distribución (Manual)", f.distribution_paragraph),
                ("Altitud", elev),
                ("Vertientes", ", ".join(f.vertientes or []) or "no indicadas"),
                ("Regiones según el Manual", ", ".join(reg["stated"]) or "ninguna extraída"),
                ("Regiones añadidas por registros", ", ".join(reg["inferred"]) or "ninguna"),
                ("Floración", ", ".join(MESES[m - 1] for m in (f.flowering_months or []))
                 or "no indicada"),
                ("Forma de crecimiento", ", ".join(f.habits or []) or "no extraída"),
                ("Registros de GBIF", f"{n_snapshot} en la descarga usada por el sistema"),
            ]
            r = row + 1
            for label, value in facts:
                _cell(ws, f"A{r}", label, bold=True)
                _cell(ws, f"B{r}", value)
                ws.row_dimensions[r].height = max(15, 14 * (len(str(value)) // 95 + 1))
                r += 1
            _cell(ws, f"A{r}", "¿El mapa es consistente con lo que usted conoce de la "
                               "especie?", bold=True)
            _cell(ws, f"B{r}", "", fill=FILL)
            dv.add(f"B{r}")
            ws.row_dimensions[r].height = 30
            _cell(ws, f"A{r + 1}", "Comentario sobre el mapa", bold=True)
            _cell(ws, f"B{r + 1}", "", fill=FILL)
            ws.row_dimensions[r + 1].height = 45
            r += 2
            try:
                m = render(sp)
                path = (m.get("value") or {}).get("map_path")
            except Exception:
                path = None
            if path and Path(path).exists():
                small = tmp / (sp.replace(" ", "_") + ".png")
                im = PILImage.open(path).convert("RGB")
                w = 560
                im = im.resize((w, int(im.height * w / im.width)))
                im.save(small, optimize=True)
                img = XLImage(str(small))
                ws.add_image(img, f"B{r}")
                rows_needed = int(im.height / 20) + 2
            else:
                _cell(ws, f"B{r}", "No se pudo generar el mapa de esta especie.")
                rows_needed = 2
            row = r + rows_needed + 1
    conn.close()
    return anchors


def build() -> Path:
    items = select()
    species = sorted({it["q"]["species"] for it in items})
    wb = Workbook()
    _instructions(wb.active)
    review = wb.create_sheet("Revisión")
    fichas = wb.create_sheet("Fichas y mapas")
    with tempfile.TemporaryDirectory() as tmp:
        anchors = _fichas(fichas, species, Path(tmp))
        _review(review, items, anchors)
        out = DATA / "expert_review_tramo1.xlsx"
        wb.save(out)

    verdicts = {(v["qid"], v["arm"]): v["direct"] for v in _jsonl("verdicts_main.jsonl")
                if not v.get("error")}
    with io.open(DATA / "expert_review_key.jsonl", "w", encoding="utf-8") as fh:
        for i, it in enumerate(items, 1):
            q = it["q"]
            fh.write(json.dumps({
                "n": i, "qid": q["qid"], "qtype": q["qtype"], "species": q["species"],
                "A": it["A"], "B": it["B"],
                "checker_A": verdicts.get((q["qid"], it["A"])),
                "checker_B": verdicts.get((q["qid"], it["B"]))}, ensure_ascii=False) + "\n")
    print(f"{len(items)} questions over {len(species)} species -> {out} "
          f"({out.stat().st_size // 1024} KB)")
    return out


if __name__ == "__main__":
    build()
